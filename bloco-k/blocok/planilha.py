"""Abre a planilha de estoque, qualquer que seja o formato.

.xlsx e .xlsm são lidos direto e .csv é lido como texto. Os demais (.xls, .xlsb...) são
abertos pelo Excel instalado, somente leitura, e salvos como uma cópia .xlsx temporária
em %TEMP%. O arquivo original nunca é alterado.
"""
from __future__ import annotations

import csv
import io
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .nucleo import ErroLeitura

LIDOS_DIRETO = {".xlsx", ".xlsm"}
SCRIPT_EXCEL = Path(__file__).with_name("converter_excel.ps1")


@dataclass
class Aba:
    nome: str
    linhas: list[list]   # linhas[0] é a linha 1 do Excel; linhas[i][0] é a coluna A
    abas: list[str]      # nomes de todas as abas do arquivo, na ordem


def ler_aba(caminho, aba: int | str = -1) -> Aba:
    """Lê uma aba (pelo índice, -1 = última, ou pelo nome) e devolve os valores das células."""
    return _abrir(caminho, lambda livro: _uma(livro, aba))


def ler_abas(caminho) -> list[Aba]:
    """Lê todas as abas do arquivo, na ordem. Bom para planilhas com poucas abas."""
    return _abrir(caminho, _todas)


# ---------------------------------------------------------------- internos

def _abrir(caminho, ler):
    caminho = Path(caminho)
    if not caminho.exists():
        raise ErroLeitura(f"Arquivo não encontrado: {caminho}")
    sufixo = caminho.suffix.lower()
    if sufixo == ".csv":
        aba = _ler_csv(caminho)
        return aba if ler is not _todas else [aba]
    if sufixo in LIDOS_DIRETO:
        return _com_openpyxl(caminho, ler)
    with tempfile.TemporaryDirectory(prefix="blocok_") as tmp:
        copia = Path(tmp) / "copia.xlsx"
        _converter_com_excel(caminho, copia)
        return _com_openpyxl(copia, ler)


def _com_openpyxl(caminho: Path, ler):
    try:
        livro = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    except Exception as e:  # arquivo corrompido, protegido por senha etc.
        raise ErroLeitura(f"Não consegui abrir a planilha: {e}") from e
    try:
        return ler(livro)
    finally:
        livro.close()


def _valores(planilha) -> list[list]:
    # No modo read_only o openpyxl acredita na dimensão que o arquivo declara. Planilha
    # gerada por outro programa às vezes declara menos linhas do que tem, e aí a leitura
    # para no meio sem erro nenhum. reset_dimensions faz ele contar as linhas de verdade.
    if hasattr(planilha, "reset_dimensions"):
        try:
            planilha.reset_dimensions()
        except Exception:
            pass
    return [list(r) for r in planilha.iter_rows(min_row=1, values_only=True)]


def _uma(livro, aba) -> Aba:
    nomes = livro.sheetnames
    try:
        planilha = livro[nomes[aba]] if isinstance(aba, int) else livro[aba]
    except (IndexError, KeyError):
        raise ErroLeitura(f"A aba {aba!r} não existe. Abas do arquivo: {', '.join(nomes)}") from None
    return Aba(planilha.title, _valores(planilha), nomes)


def _todas(livro) -> list[Aba]:
    nomes = livro.sheetnames
    return [Aba(nome, _valores(livro[nome]), nomes) for nome in nomes]


def _ler_csv(caminho: Path) -> Aba:
    bruto = caminho.read_bytes()
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            conteudo = bruto.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    primeira = conteudo.splitlines()[0] if conteudo else ""
    separador = max((";", ",", "\t"), key=primeira.count)
    linhas = [list(r) for r in csv.reader(io.StringIO(conteudo), delimiter=separador)]
    return Aba(caminho.stem, linhas, [caminho.stem])


def _converter_com_excel(origem: Path, destino: Path) -> None:
    comando = ["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
               "-File", str(SCRIPT_EXCEL), str(origem), str(destino)]
    try:
        r = subprocess.run(comando, capture_output=True, text=True, timeout=300,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except subprocess.TimeoutExpired:
        raise ErroLeitura("O Excel demorou mais de 5 minutos para abrir o arquivo.") from None
    if r.returncode != 0 or not destino.exists():
        detalhe = (r.stderr or r.stdout or "").strip()
        raise ErroLeitura("Não consegui abrir o arquivo pelo Excel. Para arquivos que não são .xlsx nem .csv, "
                          "é preciso ter o Excel instalado; outra saída é salvar a planilha como .xlsx."
                          + (f"\n\nDetalhe: {detalhe}" if detalhe else ""))
