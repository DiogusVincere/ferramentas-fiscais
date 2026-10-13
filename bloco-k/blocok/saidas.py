"""As duas planilhas de importação, iguais para todas as empresas.

1. BLOCO K - JUNHO2026 - CADASTRAR NO DOMINIO.xlsx  (cadastro de produtos)
2. BLOCO K 062026 modelo Dominio.xlsx              (Posse de Produtos: registros 01 e 02)

Layout e formatação seguem o modelo de importação de Posse de Produtos do sistema
contábil (Domínio). No modelo, "Quantidade" é a quantidade somada do código e "Saldo"
é o valor total do código em estoque.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .nucleo import ItemConsolidado

MESES = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO",
         "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]

TEXTO = "@"
FMT_ZERO = "0.000;[Red]0.000"
FMT_DECIMAL = "0.0000000000"
FMT_DATA = "mm-dd-yy"   # data curta do Windows (aparece 30/06/2026)

CAB_CADASTRO = ["Produto", "Descrição", "Unidade de Medida", "NCM", "CEST", "Codigo da Conta Contabil",
                "Quantidade", "Valor Unitario", "Total", "ICMS Recuperável"]
LARG_CADASTRO = [19.3, 59.9, 21.3, 9.0, 8.1, 27.0, 13.9, 16.0, 10.6, 20.3]

CAB_MODELO = [
    "Tipo 02",
    "Código Produto",
    'Indicador de Propriedade \n("0"-Do informante em seu poder / \n"1"-Do informante em posse de terceiros / \n'
    '"2"-De terceiros em posse do informante)',
    'Cliente ou Fornecedor ("C"=CLiente  /  "F"=Fornecedor)',
    "Inscrição Participante (CNPJ)",
    "Quantidade",
    "Saldo",
    "Código da Conta Contábil",
]
LARG_MODELO = [10.3, 35.6, 39.4, 18.1, 19.7, 17.9, 13.0, 9.7]


def nomes_arquivos(competencia: dt.date, sufixo: str = "") -> tuple[str, str]:
    """Nomes do cadastro e do modelo. sufixo separa importações da mesma empresa (ex.: " - FILIAL")."""
    mes, ano = competencia.month, competencia.year
    return (f"BLOCO K - {MESES[mes - 1]}{ano} - CADASTRAR NO DOMINIO{sufixo}.xlsx",
            f"BLOCO K {mes:02d}{ano} modelo Dominio{sufixo}.xlsx")


def _conta_numero(conta: str):
    return int(conta) if str(conta).isdigit() else conta


def _larguras(ws, larguras):
    for i, largura in enumerate(larguras):
        ws.column_dimensions[chr(ord("A") + i)].width = largura


def gerar_cadastro(itens: list[ItemConsolidado], conta: str, destino: Path) -> Path:
    """Um código por linha; quantidade, valor e total zerados, como no exemplo."""
    livro = Workbook()
    ws = livro.active
    ws.title = "Planilha1"
    arial = Font(name="Arial", size=10)
    cab_fonte = Font(name="Arial", size=10, bold=True)
    cab_fundo = PatternFill("solid", fgColor="C0C0C0")

    for col, titulo in enumerate(CAB_CADASTRO, start=1):
        c = ws.cell(1, col, titulo)
        c.font, c.fill, c.number_format = cab_fonte, cab_fundo, TEXTO

    vistos = set()
    linha = 1
    for it in itens:
        if it.codigo in vistos:          # o mesmo produto com outra posse: cadastra uma vez só
            continue
        vistos.add(it.codigo)
        linha += 1
        valores = [it.codigo, it.descricao, it.unidade, it.ncm, None, str(conta), 0, 0, 0, None]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(linha, col, v)
            c.font = arial
            c.number_format = FMT_ZERO if col in (7, 8, 9) else TEXTO

    _larguras(ws, LARG_CADASTRO)
    ws.auto_filter.ref = f"A1:J{max(linha, 2)}"
    livro.save(destino)
    return destino


def gerar_modelo(itens: list[ItemConsolidado], competencia: dt.date, conta: str, destino: Path) -> Path:
    """Registro 01 (competência) e um registro 02 por código consolidado."""
    livro = Workbook()
    ws = livro.active
    ws.title = "Sheet1"
    arial = Font(name="Arial", size=10)
    calibri_b = Font(name="Calibri", size=11, bold=True)
    esquerda = Alignment(horizontal="left")
    topo = Alignment(vertical="top")

    c = ws.cell(1, 1, "Tipo 01")
    c.font, c.alignment, c.number_format = calibri_b, esquerda, TEXTO
    c = ws.cell(1, 2, "Competencia")
    c.font, c.alignment = arial, topo
    c = ws.cell(2, 1, "01")
    c.font, c.alignment, c.number_format = arial, esquerda, TEXTO
    c = ws.cell(2, 2, dt.datetime(competencia.year, competencia.month, competencia.day))
    c.font, c.alignment, c.number_format = arial, Alignment(horizontal="left", vertical="top"), FMT_DATA
    ws.row_dimensions[1].height = 12.75
    ws.row_dimensions[2].height = 12.75

    for col, titulo in enumerate(CAB_MODELO, start=1):
        c = ws.cell(3, col, titulo)
        c.font = calibri_b
        c.alignment = Alignment(wrap_text=True, horizontal={1: "left", 2: "center", 7: "center"}.get(col))
        c.number_format = TEXTO if col == 1 else (FMT_DECIMAL if col in (6, 7) else "General")
    ws.row_dimensions[3].height = 66

    conta_valor = _conta_numero(conta)
    for i, it in enumerate(itens, start=4):
        valores = [
            ("02", TEXTO, topo),
            (it.codigo, TEXTO, None),
            (it.indicador if it.indicador is not None else 0, "General", topo),
            (it.participante or None, TEXTO, None),
            (it.cnpj_participante or None, TEXTO, None),
            (it.quantidade, FMT_DECIMAL, None),
            (it.total, FMT_DECIMAL, None),          # Saldo = valor total em estoque
            (conta_valor, "#,##0", topo),
        ]
        for col, (v, fmt, alinhamento) in enumerate(valores, start=1):
            c = ws.cell(i, col, v)
            c.font, c.number_format = arial, fmt
            if alinhamento:
                c.alignment = alinhamento

    _larguras(ws, LARG_MODELO)
    ws.auto_filter.ref = f"A3:H{max(3 + len(itens), 4)}"
    livro.save(destino)
    return destino
