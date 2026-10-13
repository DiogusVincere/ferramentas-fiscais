"""Gera exemplos/estoque-exemplo.xlsx: a planilha de estoque fictícia da Empresa Exemplo Ltda.

É o arquivo que o leitor leitores/exemplo.py sabe ler. Serve para experimentar a
ferramenta sem precisar de uma planilha de verdade.

    python gerar_planilha_exemplo.py            -> grava ao lado deste script
    python gerar_planilha_exemplo.py <destino>  -> grava no caminho indicado

A planilha traz 30 produtos inventados e alguns casos de propósito, para a aba
"Avisos" ter o que mostrar: um código repetido em duas linhas (as quantidades e os
totais são somados), um NCM que o Excel guardou como número e perdeu o zero da frente
(o leitor devolve o zero), uma linha sem o valor total (o total sai de quantidade ×
unitário), um código com mais de 14 caracteres e uma linha de total no fim (ignorada
por não ter código).
"""
import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

PASTA = Path(os.path.abspath(__file__)).parent
DESTINO_PADRAO = PASTA / "estoque-exemplo.xlsx"

# (código, descrição, unidade, NCM, quantidade, valor unitário)
PRODUTOS = [
    ("P0001", "Parafuso sextavado M6 x 40 zincado", "UN", "73181500", 1200, 0.42),
    ("P0002", "Porca sextavada M6 zincada", "UN", "73181600", 1500, 0.18),
    ("P0003", "Arruela lisa M6", "UN", "73182200", 2000, 0.07),
    ("P0004", "Chapa de aço carbono 1,5 mm 1000 x 2000", "PC", "72085200", 18, 312.50),
    ("P0005", "Chapa de aço inox 304 1,0 mm 1000 x 2000", "PC", "72193300", 6, 890.00),
    ("P0006", "Tubo de aço quadrado 30 x 30 x 1,2 mm 6 m", "PC", "73066100", 40, 98.70),
    ("P0007", "Barra chata 1 x 1/8 pol 6 m", "PC", "72166100", 25, 44.90),
    ("P0008", "Eletrodo revestido 2,5 mm (caixa 5 kg)", "CX", "83111000", 12, 86.00),
    ("P0009", "Disco de corte 7 pol", "UN", "68042290", 60, 9.80),
    ("P0010", "Disco de desbaste 7 pol", "UN", "68042290", 30, 14.60),
    ("P0011", "Tinta esmalte sintético cinza 3,6 L", "LT", "32089010", 15, 78.90),
    ("P0012", "Primer anticorrosivo 3,6 L", "LT", "32089010", 10, 92.40),
    ("P0013", "Solvente para esmalte 5 L", "LT", "38140010", 8, 61.00),
    ("P0014", "Lixa ferro grão 80", "FL", "68052000", 200, 1.95),
    ("P0015", "Lixa ferro grão 120", "FL", "68052000", 150, 1.95),
    ("P0016", "Fita isolante 19 mm x 20 m", "RL", "39191020", 48, 6.30),
    ("P0017", "Cabo flexível 2,5 mm² preto (rolo 100 m)", "RL", "85444900", 7, 189.00),
    ("P0018", "Cabo flexível 2,5 mm² azul (rolo 100 m)", "RL", "85444900", 5, 189.00),
    ("P0019", "Disjuntor monopolar 20 A", "UN", "85362000", 36, 17.80),
    ("P0020", "Tomada 2P+T 10 A branca", "UN", "85366990", 80, 8.45),
    ("P0021", "Luva de raspa cano curto (par)", "PAR", "42032900", 90, 11.20),
    ("P0022", "Óculos de proteção incolor", "UN", "90049090", 50, 7.90),
    ("P0023", "Protetor auricular tipo concha", "UN", "63079010", 25, 29.00),
    ("P0024", "Capacete de segurança amarelo", "UN", "65061000", 20, 24.50),
    ("P0025", "Óleo lubrificante 20W50 1 L", "LT", "27101932", 24, 32.00),
    ("P0026", "Café torrado e moído 500 g (copa)", "UN", "09012100", 18, 21.70),
    ("P0027", "Estopa branca (fardo 1 kg)", "KG", "63109000", 30, 12.80),
    ("P0028", "Rolamento 6204 2RS", "UN", "84821010", 40, 15.90),
    ("P0029", "Correia em V A-42", "UN", "40103100", 22, 19.60),
    ("PARAF-M8X60-INOX", "Parafuso sextavado M8 x 60 inox", "UN", "73181500", 300, 1.35),
]


def montar(destino: Path) -> Path:
    livro = Workbook()
    ws = livro.active
    ws.title = "Estoque"
    negrito = Font(name="Arial", size=10, bold=True)
    normal = Font(name="Arial", size=10)
    cinza = PatternFill("solid", fgColor="D9D9D9")

    ws.cell(1, 1, "Empresa Exemplo Ltda").font = Font(name="Arial", size=12, bold=True)
    ws.cell(2, 1, "Inventário de estoque em 31/07/2026").font = normal
    titulos = ["Código", "Descrição", "Unidade", "NCM", "Quantidade", "Valor unitário", "Valor total"]
    for col, titulo in enumerate(titulos, start=1):
        c = ws.cell(4, col, titulo)
        c.font, c.fill = negrito, cinza
        c.alignment = Alignment(horizontal="center")

    linha = 5
    for codigo, descricao, unidade, ncm, qtd, unitario in PRODUTOS:
        total = round(qtd * unitario, 2)
        if codigo == "P0026":
            ncm = int(ncm)         # guardado como número: o Excel come o zero da frente (9012100)
        if codigo == "P0013":
            total = None           # sem valor total: sai de quantidade × unitário
        valores = [codigo, descricao, unidade, ncm, qtd, unitario, total]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(linha, col, v)
            c.font = normal
            if col == 4 and isinstance(v, str):
                c.number_format = "@"
            elif col in (6, 7):
                c.number_format = "#,##0.00"
        linha += 1

    # o mesmo código numa segunda linha (outro lote): quantidades e totais são somados
    extra = ["P0004", "Chapa de aço carbono 1,5 mm 1000 x 2000", "PC", "72085200", 4, 318.00, 1272.00]
    for col, v in enumerate(extra, start=1):
        c = ws.cell(linha, col, v)
        c.font = normal
        if col == 4:
            c.number_format = "@"
        elif col in (6, 7):
            c.number_format = "#,##0.00"
    linha += 2

    ws.cell(linha, 2, "Total geral").font = negrito
    ws.cell(linha, 7, f"=SUM(G5:G{linha - 2})").font = negrito

    for letra, largura in zip("ABCDEFG", (18, 44, 9, 11, 12, 14, 14)):
        ws.column_dimensions[letra].width = largura
    ws.freeze_panes = "A5"

    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return destino


if __name__ == "__main__":
    alvo = Path(sys.argv[1]) if len(sys.argv) > 1 else DESTINO_PADRAO
    print("Gravado:", montar(alvo))
