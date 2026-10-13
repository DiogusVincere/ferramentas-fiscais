"""Excel do resultado da comparação: resumo, todas as notas e uma aba por situação.

As cores são as mesmas de hub/tema.py (em ARGB), repetidas aqui para o módulo
continuar puro, sem depender do Tkinter.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .comparar import Resultado, reais

ESCURO = "FF1F2A37"
ESCURO_TEXTO = "FFEAEEF3"
ZEBRA = "FFF4F5F7"
VERMELHO_T = "FFF7E2DF"
AMBAR_T = "FFFBF0D9"

FMT_REAIS = 'R$ #,##0.00'
FMT_DATA = "DD/MM/YYYY"

COLUNAS = [
    ("Situação", 20), ("Nota", 14), ("Série", 8), ("Fornecedor / cliente", 34),
    ("CNPJ", 20), ("Data", 12), ("Chave da NFe", 46),
    ("Valor arquivo 1", 16), ("Valor arquivo 2", 16), ("Diferença", 14), ("Diferença %", 12),
    ("ICMS arquivo 1", 15), ("ICMS arquivo 2", 15), ("ICMS difere", 12),
    ("CFOP arquivo 1", 16), ("CFOP arquivo 2", 16), ("O que eu vi", 60),
]


def _cabecalho(ws, titulos):
    fonte = Font(name="Segoe UI", size=9, bold=True, color=ESCURO_TEXTO)
    fundo = PatternFill("solid", fgColor=ESCURO)
    for col, (titulo, largura) in enumerate(titulos, start=1):
        c = ws.cell(1, col, titulo)
        c.font, c.fill = fonte, fundo
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"


def _linha_valores(l) -> list:
    return [
        l.situacao_texto, l.numero, l.serie, l.participante,
        l.cnpj, l.data, l.chave,
        l.valor_a, l.valor_b, l.diferenca_valor, l.diferenca_percentual,
        l.icms_a, l.icms_b, "sim" if l.icms_difere else "",
        l.cfop_a, l.cfop_b,
        "; ".join(l.motivos + ([l.parecida_com] if l.parecida_com else [])),
    ]


def _aba_notas(livro, nome, linhas):
    ws = livro.create_sheet(nome[:31])
    _cabecalho(ws, COLUNAS)
    arial = Font(name="Segoe UI", size=9)
    for i, l in enumerate(linhas, start=2):
        valores = _linha_valores(l)
        destaque = None
        if l.situacao == "valor":
            destaque = PatternFill("solid", fgColor=VERMELHO_T)
        elif l.situacao in ("so_a", "so_b"):
            destaque = PatternFill("solid", fgColor=AMBAR_T)
        elif i % 2 == 0:
            destaque = PatternFill("solid", fgColor=ZEBRA)
        for col, v in enumerate(valores, start=1):
            c = ws.cell(i, col, v)
            c.font = arial
            if destaque:
                c.fill = destaque
            if col in (8, 9, 10, 12, 13):
                c.number_format = FMT_REAIS
            elif col == 11:
                c.number_format = '0.00"%"'
            elif col == 6 and isinstance(v, (dt.date, dt.datetime)):
                c.number_format = FMT_DATA
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{max(len(linhas) + 1, 2)}"
    return ws


def _aba_resumo(livro, res: Resultado):
    ws = livro.create_sheet("Resumo", 0)
    titulo = Font(name="Segoe UI", size=13, bold=True, color="FFFFFFFF")
    rotulo = Font(name="Segoe UI", size=10, bold=True)
    normal = Font(name="Segoe UI", size=10)

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 78
    c = ws.cell(1, 1, "Comparação de notas fiscais")
    c.font, c.fill = titulo, PatternFill("solid", fgColor=ESCURO)
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 26

    t = res.totais
    linhas = [
        ("", ""),
        ("Arquivo 1", res.doc_a.nome),
        ("O que entendi dele", res.doc_a.resumo()),
        ("Arquivo 2", res.doc_b.nome),
        ("O que entendi dele", res.doc_b.resumo()),
        ("Como casei as notas", res.explicacao_chave),
        ("", ""),
        ("Notas casadas", t["casadas"]),
        ("Valor confere", t["confere"]),
        ("Valor diferente", t["divergentes"]),
        ("Só no arquivo 1", t["so_a"]),
        ("Só no arquivo 2", t["so_b"]),
        ("ICMS diferente", t["icms_difere"] if res.comparou_icms else "não comparei"),
        ("CFOP diferente", t["cfop_difere"] if res.comparou_cfop else "não comparei"),
        ("", ""),
        ("Soma do arquivo 1", reais(t["soma_a"])),
        ("Soma do arquivo 2", reais(t["soma_b"])),
        ("Diferença entre as somas", reais(t["diferenca"])),
        ("Maior diferença numa nota", reais(t["maior_diferenca"])),
    ]
    linha = 2
    for a, b in linhas:
        if a:
            ws.cell(linha, 1, a).font = rotulo
            cel = ws.cell(linha, 2, b)
            cel.font = normal
            cel.alignment = Alignment(wrap_text=True, vertical="top")
        linha += 1

    if res.avisos:
        ws.cell(linha + 1, 1, "Avisos").font = rotulo
        for aviso in res.avisos:
            linha += 1
            cel = ws.cell(linha + 1, 2, aviso)
            cel.font, cel.alignment = normal, Alignment(wrap_text=True, vertical="top")
    return ws


def gerar(res: Resultado, destino) -> Path:
    """Grava o Excel do resultado e devolve o caminho."""
    destino = Path(destino)
    livro = Workbook()
    livro.remove(livro.active)

    _aba_resumo(livro, res)
    _aba_notas(livro, "Todas as notas", res.linhas)
    grupos = [
        ("Valor diferente", res.por_situacao("valor")),
        ("So no arquivo 1", res.por_situacao("so_a")),
        ("So no arquivo 2", res.por_situacao("so_b")),
        ("ICMS diferente", [l for l in res.linhas if l.icms_difere]),
        ("Confere", res.por_situacao("ok")),
    ]
    for nome, linhas in grupos:
        if linhas:
            _aba_notas(livro, nome, linhas)

    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return destino


def nome_sugerido(res: Resultado) -> str:
    competencia = res.doc_a.competencia or res.doc_b.competencia
    parte = f" {competencia[1]:02d}-{competencia[0]}" if competencia else ""
    return f"Comparacao de notas{parte}.xlsx"
